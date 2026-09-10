from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 20 * 1024 * 1024


@dataclass
class ParsedDocument:
    filename: str
    mime_type: str
    text: str
    blocks: list[dict]


def _clean(value: str | None) -> str:
    return " ".join((value or "").split())


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("İlk MVP için azami dosya boyutu 20 MB'dir.")

    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(filename, data)
    if suffix == ".docx":
        return _parse_docx(filename, data)
    raise ValueError("Desteklenmeyen dosya türü.")


def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    reader = PdfReader(BytesIO(data))
    blocks: list[dict] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = _clean(page.extract_text())
        if page_text:
            blocks.append({"location": f"Sayfa {page_number}", "text": page_text})

    text = "\n\n".join(f"[{b['location']}]\n{b['text']}" for b in blocks)
    return ParsedDocument(
        filename=filename,
        mime_type="application/pdf",
        text=text,
        blocks=blocks,
    )


def _parse_docx(filename: str, data: bytes) -> ParsedDocument:
    document = Document(BytesIO(data))
    blocks: list[dict] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        paragraph_text = _clean(paragraph.text)
        if paragraph_text:
            blocks.append({"location": f"Paragraf {index}", "text": paragraph_text})

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            rows.append(" | ".join(cells))
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text:
            blocks.append({"location": f"Tablo {table_index}", "text": table_text})

    text = "\n\n".join(f"[{b['location']}]\n{b['text']}" for b in blocks)
    return ParsedDocument(
        filename=filename,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        text=text,
        blocks=blocks,
    )
